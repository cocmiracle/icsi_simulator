"""Попълнената таблица -> lang/*.js.

  python tools/i18n/i18n_import.py  превод.xlsx  [en,de]   (по подразбиране: всички езици в таблицата)

- ID (скритата колона) се сверява с tools/i18n/keys.json и с копието в таблицата; пипнат ID се взема по №.
- {0}, {1}… и HTML знаците (<b>, <kbd>…) трябва да са същите като в ID; ако не са - редът се пропуска и се съобщава.
- Български: записва се само там, където текстът е поправен (lang/bg.js); иначе остава текстът от програмата.
"""
import json, os, re, sys
from openpyxl import load_workbook
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from i18n_export import load_lang, ROOT, NAMES

PH = lambda s: sorted(re.findall(r"\{\d+\}", s))
TAGS = lambda s: sorted(re.findall(r"</?([a-zA-Z]+)", s))

def write_lang(code, d, name):
    head = {"bg": "// Български: текстовете са в програмата. Тук са само поправките: \"текст в програмата\": \"поправен текст\"."}.get(
        code, f"// {name}. Ключ = българският текст в програмата; стойност = превод. {{0}}, {{1}}… = стойности от програмата.\n// Липсващите остават на български. Попълва се от таблицата: tools/i18n/i18n_import.py")
    lines = [head, f'I18N.add("{code}", {{'] + ["  " + json.dumps(k, ensure_ascii=False) + ": " + json.dumps(v, ensure_ascii=False) + "," for k, v in d.items()] + ["});"]
    open(os.path.join(ROOT, "lang", code + ".js"), "w", encoding="utf-8").write("\n".join(lines) + "\n")

def main():
    wb = load_workbook(sys.argv[1]); ws = wb["Превод"]
    keys = json.load(open(os.path.join(ROOT, "tools", "i18n", "keys.json"), encoding="utf-8"))
    sent = {r[0]: r[1] for r in wb["meta"].iter_rows(min_row=2, values_only=True)} if "meta" in wb.sheetnames else {}
    head = [c.value for c in ws[1]]
    cols = {re.search(r"\[(\w+)\]", h).group(1): j for j, h in enumerate(head) if h and re.search(r"\[(\w+)\]$", h)}
    only = sys.argv[2].split(",") if len(sys.argv) > 2 else list(cols)
    report, new = [], {c: {} for c in only}
    olds = {c: load_lang(c) for c in only}
    for row in ws.iter_rows(min_row=2, values_only=True):
        kid, num = row[0], row[1]
        key = kid if kid in keys else sent.get(num) or (keys[num] if isinstance(num, int) and num < len(keys) else None)
        if kid != key: report.append(f"№{num}: скритият ID е пипнат - взет е оригиналът")
        if key is None or key not in keys: report.append(f"№{num}: непознат ред - пропуснат"); continue
        for c in only:
            v = row[cols[c]]
            if v is None or str(v).strip() == "": continue
            v = str(v)                                        # без изрязване: някои текстове започват с интервал
            if PH(v) != PH(key) or TAGS(v) != TAGS(key):
                report.append(f"№{num} [{c}]: {{0}}… или <знаци> не съвпадат с оригинала - остава старият текст")
                if key in olds[c]: new[c][key] = olds[c][key]
                continue
            if c == "bg" and v == key: continue
            new[c][key] = v
    for c in only:
        old = olds[c]
        changed = sum(1 for k in set(old) | set(new[c]) if old.get(k) != new[c].get(k))
        write_lang(c, new[c], NAMES.get(c, c))
        print(f"lang/{c}.js: {len(new[c])} текста, променени {changed}")
    # списъкът с езици
    present = [f[:-3] for f in sorted(os.listdir(os.path.join(ROOT, "lang"))) if f.endswith(".js") and f != "languages.js"]
    order = ["bg", "en"] + [c for c in present if c not in ("bg", "en")]
    body = "\n".join(f'  ["{c}", "{NAMES.get(c, c)}"],' for c in order if c in present)
    open(os.path.join(ROOT, "lang", "languages.js"), "w", encoding="utf-8").write(
        "// Наличните езици: [код, име на самия език]. Нов език: файл lang/<код>.js и ред тук.\nI18N.langs = [\n" + body + "\n];\n")
    print("\n".join(report) if report else "Проверка: всичко е наред.")

if __name__ == "__main__":
    main()
