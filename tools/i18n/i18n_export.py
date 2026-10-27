"""Таблица за превод: tools/i18n/keys.json + lang/*.js -> Excel.

  python tools/i18n/i18n_export.py  превод.xlsx  [en,bg,de ...]

Колони: № | English [en] | Български [bg] | ... | Къде | Вид | Бележка; скрита и заключена колона ID
(текстът в програмата). Листът е защитен (паролата е в PASSWORD по-долу), пишат се само преводите и бележките.
"""
import json, os, re, sys, datetime
from openpyxl import Workbook
from openpyxl.workbook.protection import WorkbookProtection
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side, Protection

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
PASSWORD = "icsi"
NAMES = {"bg": "Български", "en": "English", "de": "Deutsch", "fr": "Français", "es": "Español", "it": "Italiano"}

def load_lang(code):
    p = os.path.join(ROOT, "lang", code + ".js")
    if not os.path.exists(p): return {}
    t = open(p, encoding="utf-8").read()
    body = t[t.index("{", t.index("I18N.add(")): t.rindex("}") + 1]
    body = re.sub(r",\s*}$", "}", re.sub(r"^\s*//.*$", "", body, flags=re.M).strip())
    return json.loads(body)

def main():
    out = sys.argv[1]
    codes = (sys.argv[2] if len(sys.argv) > 2 else "en,bg").split(",")
    meta = json.load(open(os.path.join(ROOT, "tools", "i18n", "keys_meta.json"), encoding="utf-8"))
    keys = json.load(open(os.path.join(ROOT, "tools", "i18n", "keys.json"), encoding="utf-8"))
    langs = {c: load_lang(c) for c in codes}
    wb = Workbook(); info = wb.active; info.title = "Указания"
    for line in open(os.path.join(ROOT, "tools", "i18n", "instructions_bg.txt"), encoding="utf-8").read().split("\n"):
        info.append([line])
    info.column_dimensions["A"].width = 120; info["A1"].font = Font(bold=True, size=14, color="1F4E79")
    for r in info.iter_rows(min_row=2): r[0].alignment = Alignment(wrap_text=True, vertical="top")
    ws = wb.create_sheet("Превод")
    head = ["ID", "№"] + [f"{NAMES.get(c, c)} [{c}]" for c in codes] + ["Къде", "Вид", "Бележка"]
    ws.append(head)
    order = sorted(range(len(keys)), key=lambda i: (meta[keys[i]]["order"], i))
    for i in order:
        k = keys[i]
        vals = [k if (c == "bg" and k not in langs["bg"]) else langs[c].get(k, "") for c in codes]
        ws.append([k, i] + vals + [meta[k]["where"], meta[k]["kind"], ""])
    ws.column_dimensions["A"].hidden = True
    widths = {"B": 6}
    for j, c in enumerate(codes): widths[chr(ord("C") + j)] = 55
    n = len(codes); widths[chr(ord("C") + n)] = 20; widths[chr(ord("C") + n + 1)] = 12; widths[chr(ord("C") + n + 2)] = 30
    for col, w in widths.items(): ws.column_dimensions[col].width = w
    fill_h = PatternFill("solid", fgColor="DCE6F0"); fill_e = PatternFill("solid", fgColor="FFF8DC"); fill_new = PatternFill("solid", fgColor="FDE2E2")
    thin = Side(style="thin", color="CCCCCC")
    for c in ws[1]: c.font = Font(bold=True); c.fill = fill_h; c.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical="top"); c.border = Border(bottom=thin)
        for j in range(n):
            cell = row[2 + j]; cell.protection = Protection(locked=False)
            cell.fill = fill_new if not cell.value else fill_e
        row[2 + n + 2].protection = Protection(locked=False); row[2 + n + 2].fill = fill_e
        if row[2 + n + 1].value == "медицински": row[2 + n + 1].font = Font(bold=True, color="9A3412")
    ws.freeze_panes = "C2"; ws.auto_filter.ref = f"B1:{chr(ord('C') + n + 2)}{ws.max_row}"
    ws.protection.sheet = True; ws.protection.password = PASSWORD
    ws.protection.autoFilter = False; ws.protection.sort = False; ws.protection.formatColumns = True
    # копие на изпратените ID по номер - за проверката при връщане
    m = wb.create_sheet("meta"); m.append(["exported", datetime.date.today().isoformat(), ",".join(codes)])
    for i, k in enumerate(keys): m.append([i, k])
    m.sheet_state = "veryHidden"
    wb.security = WorkbookProtection(workbookPassword=PASSWORD, lockStructure=True)   # листовете не се показват, местят и трият
    wb.active = 0
    wb.save(out)
    print(f"{out}: {len(keys)} реда, езици {codes}")

if __name__ == "__main__":
    main()
